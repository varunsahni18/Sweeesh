#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Sweesh investor pitch deck (16:9, .pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import os

A = "/Users/varunsahni/Desktop/sweesh/pitchdeck/assets"
OUT = "/Users/varunsahni/Desktop/sweesh/pitchdeck/Sweesh_Pitch_Deck.pptx"

# ---------- palette ----------
NAVY   = RGBColor.from_string("0A1220")
NAVY2  = RGBColor.from_string("101B33")
PANEL  = RGBColor.from_string("1B2A4A")
INK    = RGBColor.from_string("0F172A")
MUTED  = RGBColor.from_string("5B6472")
FAINT  = RGBColor.from_string("98A2B3")
LINE   = RGBColor.from_string("E5E7EB")
WHITE  = RGBColor.from_string("FFFFFF")
OFF    = RGBColor.from_string("F6F8FB")

GREEN  = RGBColor.from_string("059669")
GREEN_B= RGBColor.from_string("34D399")
GREEN_S= RGBColor.from_string("ECFDF5")
AMBER  = RGBColor.from_string("B45309")
AMBER_B= RGBColor.from_string("FBBF24")
AMBER_S= RGBColor.from_string("FFFBEB")
GRAY   = RGBColor.from_string("6B7280")
SKY    = RGBColor.from_string("38BDF8")
SKY_D  = RGBColor.from_string("0369A1")
INDIGO = RGBColor.from_string("818CF8")
PURPLE = RGBColor.from_string("C084FC")

FONT = "Helvetica Neue"
MONO = "Menlo"

EMU_IN = 914400

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ---------- helpers ----------
def slide(bg=None):
    s = prs.slides.add_slide(BLANK)
    if bg is not None:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
    return s

def _set_font(r, size, color, bold=False, italic=False, font=FONT):
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf

def para(tf, segs, align=PP_ALIGN.LEFT, sa=0, sb=0, line=1.0, level=0, first=False):
    """segs: list of (text, size, color) or (text,size,color,bold) or (text,size,color,bold,italic)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(sa)
    p.space_before = Pt(sb)
    p.line_spacing = line
    for seg in segs:
        t = seg[0]
        sz = seg[1]
        c = seg[2]
        b = seg[3] if len(seg) > 3 else False
        it = seg[4] if len(seg) > 4 else False
        r = p.add_run()
        r.text = t
        _set_font(r, sz, c, b, it)
    return p

def rect(s, x, y, w, h, fill=None, line=None, radius=None, lw=1.0):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sp = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def shape_text(sp, segs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=0.07):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(m)
    tf.margin_right = Inches(m)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    for seg in segs:
        t, sz, c = seg[0], seg[1], seg[2]
        b = seg[3] if len(seg) > 3 else False
        r = p.add_run()
        r.text = t
        _set_font(r, sz, c, b)
    return sp

def node(s, x, y, w, h, text, fill, txtc, size=11, sub=None, subc=None):
    b = rect(s, x, y, w, h, fill=fill, radius=0.16)
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    _set_font(r, size, txtc, True)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        _set_font(r2, size - 2.5, subc if subc else txtc, False)
    return b

def arrow(s, x, y, w, h, fill):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = fill
    a.line.fill.background()
    a.shadow.inherit = False
    return a

def pic(s, path, x, y, max_w=None, max_h=None, align_right=False, align_bottom=False, frame=None, center=False):
    im = Image.open(path)
    iw, ih = im.size
    ar = iw / ih
    if max_w and max_h:
        w = max_w
        h = w / ar
        if h > max_h:
            h = max_h
            w = h * ar
    elif max_w:
        w = max_w
        h = w / ar
    else:
        h = max_h
        w = h * ar
    if center:
        left = (13.333 - w) / 2
    else:
        left = x if not align_right else x - w
    top = y if not align_bottom else y - h
    if frame is not None:
        rect(s, left - 0.06, top - 0.06, w + 0.12, h + 0.12, fill=WHITE, line=frame, radius=0.01, lw=0.75)
    s.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))
    return w, h

def header(s, kicker, title, sub=None, dark=False, title_size=28):
    kc = GREEN_B if dark else GREEN
    tc = WHITE if dark else INK
    sc = FAINT if dark else MUTED
    para(tb(s, 0.6, 0.45, 12.1, 0.3), [(kicker, 11.5, kc, True)], first=True)
    para(tb(s, 0.6, 0.72, 12.1, 1.0), [(title, title_size, tc, True)], first=True, line=1.02)
    if sub:
        para(tb(s, 0.6, 1.56, 12.1, 0.5), [(sub, 14, sc)], first=True)

def footer(s, n, dark=False):
    c = FAINT if dark else MUTED
    para(tb(s, 0.6, 7.08, 6.0, 0.3), [("Sweesh — Investor Pitch", 9, c)], first=True)
    para(tb(s, 11.9, 7.08, 0.85, 0.3), [(str(n), 9, c)], first=True, align=PP_ALIGN.RIGHT)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

# =========================================================
# SLIDE 1 — COVER
# =========================================================
s = slide(NAVY)
rect(s, 0.6, 1.28, 0.9, 0.075, fill=GREEN_B, radius=0.5)
para(tb(s, 0.6, 1.55, 12.0, 0.5), [("SWEESH", 17, GREEN_B, True)], first=True)
para(tb(s, 0.6, 2.25, 12.1, 2.6),
     [("Every line of a resume,", 54, WHITE, True)], first=True, line=1.02)
para(tb(s, 0.6, 3.15, 12.1, 1.4), [("checkable.", 54, GREEN_B, True)], first=True, line=1.02)
para(tb(s, 0.6, 4.35, 11.2, 1.6),
     [("Sweesh is a chain of custody for facts about a person — portable as one badge, "
       "verifiable end to end, revealing nothing beyond the answer.", 17, FAINT)], first=True, line=1.3)
para(tb(s, 0.6, 6.45, 12.0, 0.4),
     [("Investor presentation  ·  2026", 12, FAINT)], first=True)
notes(s, "One-sentence thesis. Lead with the badge idea; do not explain crypto here.")

# =========================================================
# SLIDE 2 — THE TRADITIONAL RESUME
# =========================================================
s = slide(WHITE)
header(s, "THE PROBLEM — WHAT WE HAVE TODAY", "Today, a career is an unverified document.",
       "An employer takes it on faith — or pays a screening firm to phone it in by hand.")
# left column
lx, lw = 0.6, 5.3
bullets = [
    ("Nothing here is signed.", "Every line is self-asserted. No registrar, no HR system, no manager has vouched for any of it."),
    ("The proof already exists.", "It sits in a registrar's transcript, an HRIS record, an old manager's memory — disconnected, never asked."),
    ("It all forgets you.", "Each system loses interest in you the day you leave — which is exactly the day you need it to answer."),
]
y = 2.35
for title, body in bullets:
    para(tb(s, lx + 0.28, y, lw - 0.28, 0.32), [(title, 14.5, INK, True)], first=True)
    para(tb(s, lx + 0.28, y + 0.32, lw - 0.28, 0.85), [(body, 12, MUTED)], first=True, line=1.15)
    rect(s, lx, y + 0.04, 0.11, 0.11, fill=GREEN, radius=0.3)
    y += 1.42
# resume image (right)
pic(s, os.path.join(A, "resume.png"), 12.73, 1.66, max_h=5.42, align_right=True, frame=LINE)
notes(s, "Use the founder's own resume. The point is that even a strong resume is unverified.")

# =========================================================
# SLIDE 3 — HOW LONG IT TAKES
# =========================================================
s = slide(WHITE)
header(s, "THE PROBLEM — THE COST", "Verification is slow, manual, and expensive.",
       "And the cost sits on the party who captures none of the value.")
stats = [
    ("$30–100", "per candidate", "background screening"),
    ("~3 days", "of phone calls", "registrar · HR · old manager"),
    ("$20–65", "per check", "employment / income verification"),
]
x = 0.6
for big, small, cap in stats:
    card = rect(s, x, 2.4, 3.9, 2.35, fill=OFF, line=LINE, radius=0.08)
    para(tb(s, x + 0.32, 2.72, 3.3, 1.0), [(big, 40, INK, True)], first=True)
    para(tb(s, x + 0.32, 3.62, 3.3, 0.4), [(small, 14, MUTED)], first=True)
    para(tb(s, x + 0.32, 4.05, 3.3, 0.4), [(cap, 11, FAINT)], first=True)
    x += 4.11
para(tb(s, 0.6, 5.25, 12.1, 1.2),
     [("The employer wants the answer. ", 15, INK, True),
      ("The registrar or HR team produces it, gets paid nothing, and carries the liability. "
       "That asymmetry — not a missing technology — is why this has stayed broken.", 15, MUTED)],
     first=True, line=1.25)
footer(s, 3)
notes(s, "Costs from Sweesh Business Summary: screening $30-100/candidate, employment verification $20-65.")

# =========================================================
# SLIDE 4 — FOUR GAPS
# =========================================================
s = slide(WHITE)
header(s, "THE PROBLEM — WHY IT PERSISTS", "Four gaps define our opening.")
rows = [
    ("Payroll knows dates, not contribution", "Payroll systems record titles and dates, nothing about what you did.", "Manager attestations under delegated authority"),
    ("Credentials verify, issuers do not", "Blockcerts checks the document — not that the issuer is a real, accredited institution.", "Recursive authority chains to a verifier-chosen root"),
    ("Verification costs privacy", "Answering a query means revealing the record — and callbacks tell the issuer where you applied.", "Selective disclosure + zero-knowledge path proofs"),
    ("Records die with the relationship", "Your work email stops working the day you leave — the day you need it.", "Issue once at the moment of truth, prove for life"),
]
y = 2.1
for gap, why, fix in rows:
    rect(s, 0.6, y, 12.13, 1.06, fill=OFF, line=LINE, radius=0.08)
    para(tb(s, 0.9, y + 0.16, 4.3, 0.7), [(gap, 14.5, INK, True)], first=True, line=1.05)
    para(tb(s, 5.35, y + 0.2, 3.9, 0.7), [(why, 11, MUTED)], first=True, line=1.08)
    para(tb(s, 9.4, y + 0.2, 3.15, 0.7), [(fix, 11.5, GREEN, True)], first=True, line=1.08)
    y += 1.22
footer(s, 4)
notes(s, "The last row is the one to internalise: any flow that asks someone to answer later will fail.")

# =========================================================
# SLIDE 5 — THE ANSWER
# =========================================================
s = slide(NAVY)
header(s, "THE ANSWER", "No resume. One badge.", dark=True)
para(tb(s, 0.6, 2.15, 11.4, 1.4),
     [("Every fact is signed once, ", 22, WHITE, True),
      ("at the moment it becomes true", 22, GREEN_B, True),
      (" — by the only party who could have known it.", 22, WHITE, True)], first=True, line=1.25)
para(tb(s, 0.6, 3.35, 11.4, 0.9),
     [("The person collects those statements and sends one link carrying only what they choose to reveal. "
       "No callback. No phone call. No resume.", 16, FAINT)], first=True, line=1.3)
pillars = [
    ("1", "Issue at the moment of truth", "The university signs the degree when it confers it. HR signs the dates at hire and at departure."),
    ("2", "Verify the signers, not just the signature", "A signature means nothing unless you know who signed. Sweesh records authority the same way it records facts."),
    ("3", "Reveal nothing beyond the answer", "The employer learns the claim is backed by a root they trust — not who signed it, or anything not deliberately sent."),
]
x = 0.6
for num, t, b in pillars:
    card = rect(s, x, 4.7, 3.9, 2.0, fill=NAVY2, line=PANEL, radius=0.1)
    para(tb(s, x + 0.3, 4.95, 0.5, 0.5), [(num, 22, GREEN_B, True)], first=True)
    para(tb(s, x + 0.3, 5.35, 3.3, 0.5), [(t, 13.5, WHITE, True)], first=True, line=1.02)
    para(tb(s, x + 0.3, 5.82, 3.3, 0.8), [(b, 10.5, FAINT)], first=True, line=1.12)
    x += 4.11
notes(s, "The three pillars. Emphasise pillar 2 — it is the whole idea and the differentiator.")

# =========================================================
# SLIDE 6 — AUTHORITY CHAIN
# =========================================================
s = slide(WHITE)
header(s, "THE ANSWER — HOW IT WORKS", "Every claim walks to a root the verifier chose.",
       "Each arrow is an accreditation. If any link is missing, unverifiable, expired, or revoked — the claim does not pass.")

para(tb(s, 0.6, 2.12, 4.0, 0.3), [("EDUCATION", 12, SKY_D, True)], first=True)
# row A (education) — linear
eA = [("US Dept of Education", "root", INDIGO, WHITE),
      ("Accreditor", "MSCHE", SKY, WHITE),
      ("University", "New York University", SKY, WHITE),
      ("Registrar", "degree conferred", SKY, WHITE),
      ("Degree claim", "✓", GREEN, WHITE)]
x = 0.6
bw, bh, gap = 2.32, 0.95, 0.32
eA_bw, eA_gap = 2.25, 0.25
for i, (t, sub, fill, tc) in enumerate(eA):
    node(s, x, 2.5, eA_bw, bh, t, fill, tc, size=12.5, sub=sub, subc=WHITE)
    if i < len(eA) - 1:
        arrow(s, x + eA_bw, 2.5 + bh / 2 - 0.1, eA_gap, 0.2, fill)
    x += eA_bw + eA_gap

para(tb(s, 0.6, 4.1, 4.0, 0.3), [("EMPLOYMENT", 12, GREEN, True)], first=True)
# row B (employment) — with HR / manager branch
node(s, 0.6, 4.48, bw, bh, "Company register", INDIGO, WHITE, size=12.5, sub="root · Companies House / SoS", subc=WHITE)
arrow(s, 0.6 + bw, 4.48 + bh / 2 - 0.1, gap, 0.2, INDIGO)
node(s, 0.6 + bw + gap, 4.48, bw, bh, "Employer", GREEN, WHITE, size=12.5, sub="Qualcomm · delegates authority", subc=WHITE)
# branch: two boxes stacked right of employer
bx = 0.6 + (bw + gap) * 2 + 0.1
node(s, bx, 4.12, 2.4, 0.78, "HR system", SKY, WHITE, size=11.5, sub="dates · title · reporting", subc=WHITE)
node(s, bx, 5.06, 2.4, 0.78, "Manager", SKY, WHITE, size=11.5, sub="work contribution", subc=WHITE)
# arrows employer -> HR and manager
arrow(s, 0.6 + (bw + gap) * 2 - 0.28, 4.5, 0.42, 0.16, SKY)
arrow(s, 0.6 + (bw + gap) * 2 - 0.28, 5.44, 0.42, 0.16, SKY)
# converge to work claim
cx = bx + 2.4 + 0.16
node(s, cx, 4.48, 2.1, bh, "Work claim", GREEN, WHITE, size=12.5, sub="signed, private", subc=WHITE)
arrow(s, cx - 0.4, 4.6, 0.44, 0.16, GREEN)
arrow(s, cx - 0.4, 5.44, 0.44, 0.16, GREEN)

para(tb(s, 0.6, 6.35, 12.1, 0.6),
     [("Scope only ever narrows. ", 13, INK, True),
      ("HR grants the manager authority but never sees what the manager writes — so your old employer "
       "cannot read what your manager said about you.", 13, MUTED)], first=True, line=1.2)
footer(s, 6)
notes(s, "The recursive-authority walk. Two chains: education and employment. Emphasise scope narrowing.")

# =========================================================
# SLIDE 7 — THE DEMO
# =========================================================
s = slide(NAVY)
header(s, "THE DEMO", "Who verified what.", dark=True)
pic(s, os.path.join(A, "demo.png"), 0, 1.7, max_w=12.13, max_h=4.28, center=True, frame=PANEL)
# four states legend
states = [
    (GREEN_B, "Authoritative", "complete chain to a root"),
    (AMBER_B, "Corroborative", "N peers agree — a count, never a tick"),
    (GRAY, "Self-asserted", "their own statement, bound to them"),
    (FAINT, "Withheld", "exists, deliberately not disclosed"),
]
x = 0.6
for dot, t, b in states:
    rect(s, x, 6.15, 0.16, 0.16, fill=dot, radius=0.5)
    para(tb(s, x + 0.28, 6.06, 2.9, 0.3), [(t, 12, WHITE, True)], first=True)
    para(tb(s, x + 0.28, 6.32, 2.9, 0.5), [(b, 9, FAINT)], first=True, line=1.05)
    x += 3.1
footer(s, 7, dark=True)
notes(s, "Live demo. Resume at centre, people who were there on the left, institutions/registers/roots on the right.")

# =========================================================
# SLIDE 8 — POSITIONING MAP
# =========================================================
s = slide(WHITE)
header(s, "MARKET — POSITIONING", "We sit where nobody else is.",
       "Holder control & privacy on one axis, depth of career evidence on the other.")
pic(s, os.path.join(A, "positioning_map.png"), 0, 1.95, max_h=4.5, frame=LINE, center=True)
para(tb(s, 0.6, 6.5, 12.1, 0.6),
     [("The portable-trust corner. ", 13.5, INK, True),
      ("Delegated authority + private chain proof + contribution-level claims — the only player with all three.", 13.5, MUTED)],
     first=True, line=1.2)
footer(s, 8)
notes(s, "Traditional verifiers (Equifax, Truework, Checkr) sit low on privacy; credential platforms (Credly) are shallow. Sweesh is alone top-right.")

# =========================================================
# SLIDE 9 — MARKET SIZE & PLAYERS
# =========================================================
s = slide(WHITE)
header(s, "MARKET — SIZE & PLAYERS", "A large, growing market.")
# left stat block
card = rect(s, 0.6, 2.2, 4.6, 4.5, fill=NAVY, radius=0.1)
para(tb(s, 0.95, 2.55, 4.0, 1.2), [("$24B+", 44, WHITE, True)], first=True)
para(tb(s, 0.95, 3.5, 4.0, 0.5), [("adjacent market today", 13, FAINT)], first=True)
para(tb(s, 0.95, 3.95, 4.0, 0.8), [("Background screening $14B + digital identity $10B", 11.5, FAINT)], first=True, line=1.2)
rect(s, 0.95, 4.7, 3.9, 0.02, fill=PANEL)
para(tb(s, 0.95, 4.95, 4.0, 1.0), [("$43–48B", 28, GREEN_B, True)], first=True)
para(tb(s, 0.95, 5.55, 4.0, 0.5), [("by 2031", 12, FAINT)], first=True)
para(tb(s, 0.95, 5.9, 4.0, 0.6), [("screening +6–7%/yr · digital identity +15–18%/yr", 10.5, FAINT)], first=True, line=1.15)
# right table
tx = 5.55
tw = 7.18
cols = [1.7, 2.15, 1.85, 1.48]
hdr = ["Group", "Players", "Provide", "Pay"]
data = [
    ["Employment verification", "Equifax, Truework", "Dates, title, income", "$20–65 / check"],
    ["Background screening", "First Advantage, Checkr", "Criminal, employment, edu", "$30–100 / candidate"],
    ["Digital credentials", "Credly, Accredible", "Diplomas, badges", "Issuer subscription"],
    ["Career networks", "Velocity Network", "Career wallet", "Network / tx fees"],
    ["Credential infra", "SpruceID, Truvera", "Issuance, wallets, APIs", "Licence + API"],
]
# header
hx = tx
for i, cw in enumerate(cols):
    rect(s, hx, 2.2, cw - 0.06, 0.44, fill=GREEN, radius=0.05)
    para(tb(s, hx + 0.08, 2.28, cw - 0.2, 0.3), [(hdr[i], 11, WHITE, True)], first=True)
    hx += cw
# rows
ry = 2.72
for row in data:
    hx = tx
    for i, cw in enumerate(cols):
        fillc = WHITE if (ry // 0.52) % 2 == 0 else OFF
        rect(s, hx, ry, cw - 0.06, 0.52, fill=fillc, line=LINE, radius=0.04, lw=0.5)
        bold = (i == 0)
        para(tb(s, hx + 0.08, ry + 0.08, cw - 0.2, 0.36),
             [(row[i], 9.5, INK if bold else MUTED, bold)], first=True, line=1.0)
        hx += cw
    ry += 0.52
footer(s, 9)
notes(s, "Source: First Advantage market presentation, via Sweesh Business Summary.")

# =========================================================
# SLIDE 10 — REVENUE PATH
# =========================================================
s = slide(WHITE)
header(s, "MARKET — REVENUE PATH", "From hackathon demo to $13.2M.")
pic(s, os.path.join(A, "revenue_path.png"), 0, 1.9, max_w=12.13, max_h=4.62, frame=LINE, center=True)
para(tb(s, 0.6, 6.62, 12.1, 0.4),
     [("Base case: ", 12, INK, True),
      ("$13M in year five (450 issuers, 200 employer/API customers) · indicative valuation $65–105M. "
       "Achievable only with a distribution partner — ATS, HRIS, or staffing platform.", 12, MUTED)],
     first=True, line=1.15)
footer(s, 10)
notes(s, "Conservative $4M / base $13M / strong $34M year-five scenarios. Valuation 6-10x ARR declining to 5-8x.")

# =========================================================
# SLIDE 11 — BUSINESS MODEL
# =========================================================
s = slide(WHITE)
header(s, "BUSINESS MODEL", "Free for candidates. Issuers and verifiers pay.")
pricing = [
    ("Basic credential issuer", "Free – $1.5k / yr"),
    ("HR / university issuer platform", "$6k – $15k / yr"),
    ("Enterprise issuer · HRIS / SIS integration", "$25k – $75k / yr"),
    ("Employer verification platform", "$12k – $60k / yr"),
    ("Individual badge verification", "$2 – $8 / badge"),
    ("ATS / background-screening API partner", "$50k – $250k / yr + usage"),
]
y = 2.15
for name, price in pricing:
    rect(s, 0.6, y, 12.13, 0.6, fill=OFF, line=LINE, radius=0.07)
    para(tb(s, 0.9, y + 0.13, 8.5, 0.34), [(name, 13, INK)], first=True)
    para(tb(s, 9.3, y + 0.13, 3.2, 0.34), [(price, 13, GREEN, True)], first=True, align=PP_ALIGN.RIGHT)
    y += 0.72
para(tb(s, 0.6, 6.55, 12.1, 0.5),
     [("Recurring subscriptions + verification fees. ", 12.5, INK, True),
      ("Not a pure per-check model — reusable credentials naturally reduce repeated checks.", 12.5, MUTED)],
     first=True, line=1.15)
footer(s, 11)
notes(s, "Closest competitor is Velocity Network; our differentiator is verifying delegated authority + private chain proofs.")

# =========================================================
# SLIDE 12 — WHY WE WIN
# =========================================================
s = slide(WHITE)
header(s, "WHY WE WIN", "We don't out-cover the incumbent. We out-scope it.")
# incumbent panel
rect(s, 0.6, 2.15, 5.85, 1.7, fill=OFF, line=LINE, radius=0.08)
para(tb(s, 0.9, 2.35, 5.3, 0.4), [("The incumbent: Equifax Work Number", 13, INK, True)], first=True)
para(tb(s, 0.9, 2.78, 5.3, 0.9),
     [("839M+ employee records, fed automatically through payroll, access sold to verifiers. "
       "A copy of most of the American workforce's history.", 11.5, MUTED)], first=True, line=1.18)
# what payroll can't do
rect(s, 6.7, 2.15, 6.03, 1.7, fill=GREEN_S, line=GREEN, radius=0.08)
para(tb(s, 7.0, 2.35, 5.4, 0.4), [("What payroll structurally cannot do", 13, GREEN, True)], first=True)
para(tb(s, 7.0, 2.78, 5.4, 0.9),
     [("Granular work claims · candidate-controlled disclosure · cross-border · verification without paying per check.", 11.5, INK)],
     first=True, line=1.18)
# cold start
para(tb(s, 0.6, 4.25, 12.1, 0.4), [("Cold start, answered.", 15, INK, True)], first=True)
para(tb(s, 0.6, 4.72, 12.1, 1.0),
     [("Sell verification-request deflection to registrars and HR teams who field these requests for free today. "
       "Credentials enter the network as exhaust — the party who acts first finally captures value.", 14, MUTED)],
     first=True, line=1.3)
# risk callout
rect(s, 0.6, 5.75, 12.13, 1.0, fill=AMBER_S, line=AMBER_B, radius=0.08)
para(tb(s, 0.9, 5.95, 11.6, 0.7),
     [("The honest risk: ", 12.5, AMBER, True),
      ("most employers restrict manager references to dates and title. Our design makes evaluation "
       "structurally unrepresentable — factual, scoped attestations with the company as legal attester.", 12.5, MUTED)],
     first=True, line=1.2)
footer(s, 12)
notes(s, "Differentiation is scope, not coverage. Acknowledge the manager-reference risk up front — it reads as competence.")

# =========================================================
# SLIDE 13 — ROADMAP & TEAM
# =========================================================
s = slide(WHITE)
header(s, "ROADMAP & TEAM", "Six weeks to demo. Then the metric that matters.")
# timeline
tl = [
    ("Week 2", "Integration, end to end, with stubs"),
    ("Week 6", "Demo: live resume → chain → ZK proof"),
    ("Month 3", "First real issuer on an HRIS / SIS sandbox"),
    ("Month 6", "Accredited issuers × auto-minted attestations"),
]
x = 0.6
for i, (when, what) in enumerate(tl):
    rect(s, x, 2.3, 2.85, 1.5, fill=OFF, line=LINE, radius=0.08)
    para(tb(s, x + 0.22, 2.5, 2.4, 0.35), [(when, 13, GREEN, True)], first=True)
    para(tb(s, x + 0.22, 2.9, 2.4, 0.8), [(what, 10.5, MUTED)], first=True, line=1.1)
    if i < 3:
        arrow(s, x + 2.85, 2.95, 0.24, 0.18, GREEN)
    x += 3.09
para(tb(s, 0.6, 4.1, 12.1, 0.8),
     [("The metric that matters: ", 14, INK, True),
      ("accredited issuers × attestations auto-minted per month with zero human touch — "
       "and whether that number grows without sales touching it.", 14, MUTED)], first=True, line=1.25)
# teams
teams = [
    ("Team A — Protocol & core", "3–4 engineers · cryptography", "Attestations, delegated authority, ZK chain proofs."),
    ("Team B — Data & graph", "2–3 · data science", "Real institution graph from ORCID, ROR, DAPIP; edge weighting."),
    ("Team C — Atlas & demo", "2–3 · frontend / WebGL", "The living map + badge inspector that carry the pitch."),
]
x = 0.6
for t, sz, b in teams:
    card = rect(s, x, 5.1, 3.9, 1.7, fill=NAVY, radius=0.09)
    para(tb(s, x + 0.28, 5.32, 3.35, 0.4), [(t, 13, WHITE, True)], first=True)
    para(tb(s, x + 0.28, 5.72, 3.35, 0.35), [(sz, 10.5, GREEN_B, True)], first=True)
    para(tb(s, x + 0.28, 6.05, 3.35, 0.6), [(b, 10, FAINT)], first=True, line=1.1)
    x += 4.11
footer(s, 13)
notes(s, "Integration checkpoint is end of week 2, not week 5. The adversarial suite (12 vectors) is the credibility slide.")

# =========================================================
# SLIDE 14 — CLOSE
# =========================================================
s = slide(NAVY)
rect(s, 0.6, 2.95, 0.9, 0.075, fill=GREEN_B, radius=0.5)
para(tb(s, 0.6, 3.25, 12.0, 0.4), [("SWEESH", 15, GREEN_B, True)], first=True)
para(tb(s, 0.6, 3.7, 12.1, 1.6),
     [("Every line of a resume, checkable.", 42, WHITE, True)], first=True, line=1.05)
para(tb(s, 0.6, 5.05, 11.4, 1.0),
     [("One badge. Verifiable end to end. Revealing nothing beyond the answer.", 18, FAINT)], first=True, line=1.3)
para(tb(s, 0.6, 6.45, 12.0, 0.4),
     [("Varun Sahni  ·  vs3606@nyu.edu", 12.5, FAINT)], first=True)
notes(s, "End on the one-liner. Open the floor.")

prs.save(OUT)
print("SAVED", OUT, "slides:", len(prs.slides._sldIdLst))
